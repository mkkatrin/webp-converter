#!/usr/bin/env python3
"""
Конвертер — картинки и видео в WebP, сжатие PDF и PDF → PowerPoint. Окно в браузере.

Запуск:  python3 webp_app.py            (или значок «Конвертер» на macOS)
         python3 webp_app.py --install  (создать приложение «Конвертер»)
Откроется страница в браузере: перетащите туда картинки или видео.
Готовые файлы сохраняются в папку «Загрузки/Конвертер».

Нужно: Pillow, pikepdf (ставится сам), для видео — ffmpeg.
Необязательно: pillow-heif — для фото iPhone в формате HEIC.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse

VERSION = "1.5"
REPO_RAW = os.environ.get("WEBP_REPO", "https://raw.githubusercontent.com/mkkatrin/webp-converter/main/")
APP_FILE = Path(__file__).resolve()

OUT_DIR = Path.home() / "Downloads" / "Конвертер"
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".heic", ".heif", ".avif", ".ico", ".webp"}
VIDEO_EXT = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v", ".wmv", ".flv", ".mpg", ".mpeg", ".3gp", ".gif"}

try:
    from pillow_heif import register_heif_opener  # type: ignore
    register_heif_opener()
except ImportError:
    pass


def find_ffmpeg():
    p = shutil.which("ffmpeg")
    if p:
        return p
    for c in ("/opt/homebrew/bin/ffmpeg", "/usr/local/bin/ffmpeg"):
        if os.path.exists(c):
            return c
    return None


def convert_image(src, dst, quality, lossless, width):
    try:
        from PIL import Image, ImageOps
    except ImportError:
        raise RuntimeError("Не установлен Pillow (pip install pillow)")
    with Image.open(src) as im:
        im = ImageOps.exif_transpose(im)
        if width and im.width > width:
            im = im.resize((width, round(im.height * width / im.width)), Image.LANCZOS)
        if im.mode not in ("RGB", "RGBA"):
            im = im.convert("RGBA" if "A" in im.getbands() or im.mode == "P" else "RGB")
        im.save(dst, "WEBP", quality=quality, lossless=lossless, method=6)


def convert_video(src, dst, quality, lossless, width, fps, is_gif):
    ff = find_ffmpeg()
    if not ff:
        raise RuntimeError("ffmpeg не найден — установите: brew install ffmpeg")
    filters = [] if is_gif else [f"fps={fps}"]
    if width:
        filters.append(f"scale='min({width},iw)':-2:flags=lanczos")
    cmd = [ff, "-y", "-loglevel", "error", "-i", str(src)]
    if filters:
        cmd += ["-vf", ",".join(filters)]
    cmd += ["-an", "-c:v", "libwebp", "-loop", "0", "-compression_level", "6"]
    cmd += ["-lossless", "1"] if lossless else ["-q:v", str(quality)]
    cmd.append(str(dst))
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError(r.stderr.strip()[-300:] or "ошибка ffmpeg")


def unique(path: Path) -> Path:
    if not path.exists():
        return path
    i = 1
    while True:
        p = path.with_name(f"{path.stem} ({i}){path.suffix}")
        if not p.exists():
            return p
        i += 1


# ---------- Сжатие PDF ----------
# Уровни: (макс. dpi, макс. пикселей по длинной стороне страницы, качество JPEG, подпись).
# Пиксельный предел нужен для больших страниц (слайды, плакаты), где dpi сам по себе
# даёт огромные картинки. 0 — без потерь.
PDF_LEVELS = [
    (None, None, None, "без потерь"),
    (200, 5000, 85, "почти без потерь"),
    (150, 3600, 78, "хорошее качество"),
    (110, 2600, 68, "заметное сжатие"),
    (85, 1920, 58, "сильное сжатие"),
    (72, 1400, 45, "очень сильное сжатие"),
]
PDF_STRONG = 4          # с этого уровня считаем потерю качества сильной
PDF_GOOD = 2            # уровень, который стараемся сохранить при разделении
PDF_STATE = {"status": "checking", "error": ""}
JOBS = {}
WORK_DIR = Path(tempfile.mkdtemp(prefix="converter-"))


PDF_DEPS = [("pikepdf", "pikepdf"), ("pypdfium2", "pypdfium2>=5,<6"), ("pptx", "python-pptx>=1,<2")]


def ensure_pdf_lib():
    """Ставит модули для PDF и PowerPoint в окружение приложения, если их ещё нет (один раз)."""
    import importlib
    missing = []
    for mod, spec in PDF_DEPS:
        try:
            importlib.import_module(mod)
        except ImportError:
            missing.append(spec)
    if not missing:
        PDF_STATE["status"] = "ready"
        return
    PDF_STATE["status"] = "installing"
    r = subprocess.run([sys.executable, "-m", "pip", "install", "-q", "--disable-pip-version-check"] + missing,
                       capture_output=True, text=True)
    if r.returncode == 0:
        import importlib
        importlib.invalidate_caches()
        PDF_STATE["status"] = "ready"
    else:
        PDF_STATE["status"] = "error"
        PDF_STATE["error"] = (r.stderr or r.stdout).strip()[-300:]


def wait_pdf_lib(timeout=600):
    t0 = time.time()
    while PDF_STATE["status"] in ("checking", "installing") and time.time() - t0 < timeout:
        time.sleep(0.5)
    if PDF_STATE["status"] != "ready":
        raise RuntimeError("модуль для PDF не установился. Проверьте интернет и перезапустите Конвертер. "
                           + PDF_STATE.get("error", ""))


def _walk_images(res, out, seen_forms, depth=0):
    """Все картинки в ресурсах, включая вложенные блоки (Form XObject) и маски прозрачности."""
    import pikepdf
    if depth > 12 or not isinstance(res, pikepdf.Dictionary):
        return
    xobjs = res.get("/XObject")
    if not isinstance(xobjs, pikepdf.Dictionary):
        return
    for _, xo in xobjs.items():
        if not isinstance(xo, pikepdf.Stream):
            continue
        st = xo.get("/Subtype")
        if st == "/Image":
            out.append(xo)
            sm = xo.get("/SMask")
            if isinstance(sm, pikepdf.Stream):
                out.append(sm)
        elif st == "/Form" and xo.objgen not in seen_forms:
            seen_forms.add(xo.objgen)
            _walk_images(xo.get("/Resources"), out, seen_forms, depth + 1)


def _page_images(page):
    out = []
    try:
        _walk_images(page.obj.get("/Resources"), out, set())
    except Exception:
        pass
    return out


def _image_weights(pdf):
    """Для каждой картинки — самая большая страница, на которой она стоит (в дюймах),
    и «вес» каждой страницы (сколько байт картинок на ней) для деления на части."""
    pages_in = {}
    weights = []
    for page in pdf.pages:
        try:
            box = page.mediabox
            side = max(abs(float(box[2]) - float(box[0])), abs(float(box[3]) - float(box[1]))) / 72
        except Exception:
            side = 11.7
        side = side or 11.7
        wsum, counted = 0, set()
        for raw in _page_images(page):
            key = raw.objgen
            pages_in[key] = max(pages_in.get(key, 0), side)
            if key not in counted:
                counted.add(key)
                try:
                    wsum += int(raw.get("/Length", 0))
                except Exception:
                    pass
        weights.append(wsum + 5000)
    return pages_in, weights


def _recompress_images(pdf, max_dpi, cap_px, quality, mask_jpeg):
    import io
    import zlib
    import pikepdf
    from pikepdf import Name, PdfImage
    from PIL import Image
    pages_in, _ = _image_weights(pdf)
    default_side = max(pages_in.values()) if pages_in else 11.7
    # какие картинки — маски прозрачности (их сжимаем без JPEG-артефактов на краях)
    # берём только картинки, которые реально стоят на страницах (с вложенными блоками)
    targets, masks = {}, set()
    for page in pdf.pages:
        for img in _page_images(page):
            targets.setdefault(img.objgen, img)
            sm = img.get("/SMask")
            if isinstance(sm, pikepdf.Stream):
                masks.add(sm.objgen)
    for raw in targets.values():
        try:
            if raw.get("/ImageMask", False) or int(raw.get("/BitsPerComponent", 8)) < 8:
                continue  # чёрно-белые маски и сканы 1-бит не трогаем
            old_len = int(raw.get("/Length", 0))
            if old_len < 15000:
                continue
            is_mask = raw.objgen in masks
            side_in = pages_in.get(raw.objgen, default_side)
            max_px = min(int(max_dpi * side_in), cap_px)
            pil = PdfImage(raw).as_pil_image()
            if is_mask:
                pil = pil.convert("L")
            elif pil.mode in ("L", "LA", "I", "I;16", "1"):
                pil = pil.convert("L")
            else:
                pil = pil.convert("RGB")
            if max(pil.size) > max_px:
                k = max_px / max(pil.size)
                pil = pil.resize((max(1, round(pil.width * k)), max(1, round(pil.height * k))), Image.LANCZOS)
            if is_mask and not mask_jpeg:
                data, flt = zlib.compress(pil.tobytes(), 9), Name.FlateDecode
            elif is_mask:
                buf = io.BytesIO()
                pil.save(buf, "JPEG", quality=min(95, quality + 12))
                data, flt = buf.getvalue(), Name.DCTDecode
            else:
                buf = io.BytesIO()
                pil.save(buf, "JPEG", quality=quality, optimize=True, progressive=True)
                data, flt = buf.getvalue(), Name.DCTDecode
            if len(data) >= old_len * 0.95:
                continue  # не стало меньше — оставляем как было
            # цветовой профиль сохраняем, если он подходит по числу каналов
            cs = raw.get("/ColorSpace")
            n_out = 1 if pil.mode == "L" else 3
            keep_cs = (isinstance(cs, pikepdf.Array) and len(cs) == 2 and cs[0] == "/ICCBased"
                       and int(cs[1].get("/N", 0)) == n_out)
            raw.write(data, filter=flt)
            raw.Width, raw.Height = pil.width, pil.height
            if not keep_cs:
                raw.ColorSpace = Name.DeviceGray if n_out == 1 else Name.DeviceRGB
            raw.BitsPerComponent = 8
            for k in ("/DecodeParms", "/Decode"):
                if k in raw:
                    del raw[k]
            if "/Mask" in raw and isinstance(raw.Mask, pikepdf.Array):
                del raw["/Mask"]
        except Exception:
            continue  # необычную картинку оставляем как есть


def pdf_at_level(src, dst, level, pages=None):
    """Сохраняет src (или только страницы pages) в dst со сжатием уровня level."""
    import pikepdf
    with pikepdf.open(src) as pdf:
        if pages is not None:
            keep = set(pages)
            for i in range(len(pdf.pages) - 1, -1, -1):
                if i not in keep:
                    del pdf.pages[i]
        dpi, cap, q, _ = PDF_LEVELS[level]
        if dpi:
            _recompress_images(pdf, dpi, cap, q, mask_jpeg=level >= 2)
        try:
            pdf.remove_unreferenced_resources()
        except Exception:
            pass
        pdf.save(dst, compress_streams=True, recompress_flate=True,
                 object_stream_mode=pikepdf.ObjectStreamMode.generate)
    return Path(dst).stat().st_size


def pdf_fit(src, dst, target, pages=None, start=0):
    """Подбирает самый мягкий уровень, при котором файл влезает в target.
    Сначала без потерь; если не влезло — пробует средний уровень и идёт
    к более мягким (пока влезает) или к более сильным (пока не влезет).
    Возвращает (уровень, размер, {уровень: размер})."""
    sizes, files = {}, {}
    last = len(PDF_LEVELS) - 1

    def attempt(lvl):
        tmp = Path(str(dst) + f".l{lvl}")
        sizes[lvl] = pdf_at_level(src, tmp, lvl, pages)
        files[lvl] = tmp
        return sizes[lvl] <= target

    if not attempt(0):
        mid = 3
        if attempt(mid):
            lvl = mid
            while lvl > 1 and attempt(lvl - 1):
                lvl -= 1
        else:
            lvl = mid
            while lvl < last and not attempt(lvl + 1):
                lvl += 1
            lvl = min(lvl + 1, last)
    else:
        lvl = 0
    os.replace(files[lvl], dst)
    for k, f in files.items():
        if k != lvl and f.exists():
            f.unlink()
    return lvl, sizes[lvl], sizes


def split_pages(weights, parts):
    """Делит страницы на parts кусков подряд, примерно равных по «весу»."""
    n = len(weights)
    parts = max(1, min(parts, n))
    total = sum(weights)
    groups, cur, acc, left = [], [], 0, parts
    for i, w in enumerate(weights):
        cur.append(i)
        acc += w
        remaining_pages = n - i - 1
        if left > 1 and (acc >= total / parts * (len(groups) + 1) - w / 2 or remaining_pages == left - 1):
            groups.append(cur)
            cur, left = [], left - 1
    if cur:
        groups.append(cur)
    return groups


def pdf_compress_job(src, name, target_mb):
    import pikepdf
    target = int(target_mb * 1024 * 1024)
    in_size = Path(src).stat().st_size
    try:
        with pikepdf.open(src) as pdf:
            n_pages = len(pdf.pages)
            _, weights = _image_weights(pdf)
    except pikepdf.PasswordError:
        raise RuntimeError("PDF защищён паролем — снимите пароль и попробуйте снова")
    jid = os.urandom(6).hex()
    res = WORK_DIR / f"{jid}.pdf"
    lvl, size, sizes = pdf_fit(src, res, target)
    job = {"src": str(src), "name": name, "res": str(res), "target": target,
           "weights": weights, "pages": n_pages, "level": lvl, "size": size}
    JOBS[jid] = job
    fits = size <= target
    strong = lvl >= PDF_STRONG or not fits
    parts = 0
    if strong and n_pages > 1:
        # сколько частей нужно, чтобы остаться на «хорошем качестве»
        good_size = sizes.get(PDF_GOOD) or int(sizes.get(3, size) * 1.35)
        parts = max(2, -(-int(good_size * 1.08) // target))
        parts = min(parts, n_pages, 10)
    out = {"ok": True, "kind": "pdf", "id": jid, "in": in_size, "out": size, "pages": n_pages,
           "level": PDF_LEVELS[lvl][3], "level_n": lvl, "fits": fits, "target_mb": target_mb,
           "choice": strong, "parts": parts}
    if not strong:
        out.update(pdf_keep(jid))
    return out


def _out_name(stem, suffix):
    return unique(OUT_DIR / f"{stem}{suffix}")


def pdf_keep(jid):
    job = JOBS[jid]
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    dst = _out_name(Path(job["name"]).stem, " (сжатый).pdf")
    shutil.copy2(job["res"], dst)
    return {"saved": [{"name": dst.name, "url": "/out/" + quote(dst.name), "size": dst.stat().st_size,
                       "level": PDF_LEVELS[job["level"]][3]}]}


def pdf_split(jid, parts):
    job = JOBS[jid]
    groups = split_pages(job["weights"], max(2, parts))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stem = Path(job["name"]).stem
    saved = []
    for i, pages in enumerate(groups, 1):
        tmp = WORK_DIR / f"{jid}-p{i}.pdf"
        lvl, size, _ = pdf_fit(job["src"], tmp, job["target"], pages=pages)
        dst = _out_name(stem, f" — часть {i} из {len(groups)}.pdf")
        shutil.move(str(tmp), dst)
        saved.append({"name": dst.name, "url": "/out/" + quote(dst.name), "size": size,
                      "level": PDF_LEVELS[lvl][3], "pages": f"{pages[0] + 1}–{pages[-1] + 1}",
                      "fits": size <= job["target"]})
    return {"saved": saved}


# ---------- PDF → PowerPoint ----------
def _font_family(raw_name):
    """'ABCDEF+OpenSans-SemiBold' → ('Open Sans', 'SemiBold')."""
    import re
    name = raw_name.split("+", 1)[-1]
    fam, _, style = name.partition("-")
    if "," in fam:
        fam, _, style = fam.partition(",")
    fam = re.sub(r"(PSMT|MT|PS)$", "", fam)
    fam = re.sub(r"(?<=[a-z])(?=[A-Z])|(?<=[A-Z])(?=[A-Z][a-z])", " ", fam).strip()
    return fam or "Arial", style


def _page_chars(page, c):
    """Символы страницы с координатами, шрифтом, размером и цветом (через PDFium)."""
    import ctypes
    # масштаб и шрифт каждого текстового объекта: в некоторых PDF размер шрифта = 1,
    # а настоящий размер задан масштабом объекта
    objinfo = {}
    fbuf = ctypes.create_string_buffer(256)
    for obj in page.get_objects(max_depth=20):
        if obj.type != c.FPDF_PAGEOBJ_TEXT:
            continue
        try:
            m = obj.get_matrix()
            sc = abs(m.a * m.d - m.b * m.c) ** 0.5 or 1.0
        except Exception:
            sc = 1.0
        fname = ""
        try:
            f = c.FPDFTextObj_GetFont(obj.raw)
            if f and c.FPDFFont_GetBaseFontName(f, fbuf, 256):
                fname = fbuf.value.decode("utf-8", "ignore")
        except Exception:
            pass
        objinfo[ctypes.cast(obj.raw, ctypes.c_void_p).value] = (sc, fname)
    tp = page.get_textpage()
    out = []
    buf = ctypes.create_string_buffer(256)
    flags = ctypes.c_int()
    rect = c.FS_RECTF()
    rr, gg, bb, aa = (ctypes.c_uint() for _ in range(4))
    for k in range(tp.count_chars()):
        u = c.FPDFText_GetUnicode(tp.raw, k)
        if u in (0xFFFE, 0xFFFF):
            continue
        ch = chr(u)
        gen = c.FPDFText_IsGenerated(tp.raw, k) == 1
        if gen:
            out.append({"ch": ch, "gen": True})
            continue
        obj = c.FPDFText_GetTextObject(tp.raw, k)
        key = ctypes.cast(obj, ctypes.c_void_p).value if obj else None
        c.FPDFText_GetLooseCharBox(tp.raw, k, ctypes.byref(rect))
        size = c.FPDFText_GetFontSize(tp.raw, k)
        n = c.FPDFText_GetFontInfo(tp.raw, k, buf, 256, ctypes.byref(flags))
        fname = buf.value.decode("utf-8", "ignore") if n else ""
        sc, ofont = objinfo.get(key, (1.0, ""))
        if size <= 1.5 and sc > 1.5:
            size = size * sc
        fname = fname or ofont
        weight = c.FPDFText_GetFontWeight(tp.raw, k)
        col = (255, 255, 255)
        if c.FPDFText_GetFillColor(tp.raw, k, ctypes.byref(rr), ctypes.byref(gg), ctypes.byref(bb), ctypes.byref(aa)):
            col = (rr.value, gg.value, bb.value)
        out.append({"ch": ch, "gen": False, "obj": key, "l": rect.left, "r": rect.right,
                    "t": rect.top, "b": rect.bottom, "size": size, "font": fname,
                    "weight": weight, "color": col, "italic": bool(flags.value & 64)})
    tp.close()
    return out


def _hide_text(page, c):
    """Делает прямой (не повёрнутый) видимый текст невидимым, чтобы получить фон без текста.
    Возвращает множество скрытых текстовых объектов."""
    import ctypes
    hidden = set()
    for obj in page.get_objects(max_depth=20):
        if obj.type != c.FPDF_PAGEOBJ_TEXT:
            continue
        mode = c.FPDFTextObj_GetTextRenderMode(obj.raw)
        if mode == c.FPDF_TEXTRENDERMODE_INVISIBLE:
            continue
        try:
            m = obj.get_matrix()
            a, b, cc, d = m.a, m.b, m.c, m.d
        except Exception:
            a, b, cc, d = 1, 0, 0, 1
        if abs(b) > 1e-3 or abs(cc) > 1e-3 or a <= 0 or d <= 0:
            continue  # повёрнутый или отражённый текст оставляем на фоне
        c.FPDFTextObj_SetTextRenderMode(obj.raw, c.FPDF_TEXTRENDERMODE_INVISIBLE)
        hidden.add(ctypes.cast(obj.raw, ctypes.c_void_p).value)
    if hidden:
        page.gen_content()
    return hidden


def _lines_from_chars(chars, hidden):
    """Собирает символы в строки (и куски строк, если между словами большой разрыв)."""
    lines, cur, prev = [], [], None

    def flush():
        nonlocal cur
        if any(ch["ch"].strip() for ch in cur if not ch.get("gen")):
            lines.append(cur)
        cur = []

    for ch in chars:
        if ch.get("gen"):
            if ch["ch"] in "\r\n":
                flush()
                prev = None
            elif ch["ch"] == " " and cur:
                cur.append({"ch": " ", "gen": True})
            continue
        if ch["obj"] not in hidden:
            continue
        if prev is not None:
            sz = max(prev["size"], ch["size"], 1)
            new_line = abs(ch["b"] - prev["b"]) > 0.35 * sz or ch["l"] < prev["l"] - sz
            gap = ch["l"] - prev["r"] > 2.5 * sz
            if new_line or gap:
                flush()
        cur.append(ch)
        prev = ch
    flush()
    return lines


def _runs(line):
    runs = []
    style = None
    for ch in line:
        if ch.get("gen"):
            if runs:
                runs[-1][1].append(" ")
            continue
        st = (ch["font"], round(ch["size"] * 2) / 2, ch["color"], ch["weight"] >= 600, ch["italic"])
        if st != style:
            runs.append((st, [ch["ch"]]))
            style = st
        else:
            runs[-1][1].append(ch["ch"])
    res = [(st, "".join(t)) for st, t in runs]
    if res:
        res[-1] = (res[-1][0], res[-1][1].rstrip())
    return [r for r in res if r[1]]


def pdf_to_pptx(src, dst, mode="text", progress=None):
    """mode='text' — фон картинкой + редактируемый текст; mode='image' — страницы картинками."""
    import io
    import pypdfium2 as pdfium
    import pypdfium2.raw as c
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu, Pt

    pdf = pdfium.PdfDocument(str(src))
    n = len(pdf)
    w0, h0 = pdf[0].get_size()
    EMU = 12700  # в 1 пункте
    k = min(1.0, 51206400 / (w0 * EMU), 51206400 / (h0 * EMU))  # PowerPoint: слайд не больше 56″
    k = max(k, 914400 / (min(w0, h0) * EMU)) if min(w0, h0) * EMU * k < 914400 else k
    SW, SH = int(w0 * EMU * k), int(h0 * EMU * k)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SW), Emu(SH)
    blank = prs.slide_layouts[6]
    stats = {"slides": n, "text_slides": 0, "image_slides": 0}

    for i in range(n):
        page = pdf[i]
        w, h = page.get_size()
        kk = min(SW / (w * EMU), SH / (h * EMU))  # пункты страницы → EMU слайда
        ox = (SW - w * EMU * kk) / 2
        oy = (SH - h * EMU * kk) / 2
        scale = min(220 / 72, 3000 / max(w, h))
        lines = []
        if mode == "text":
            chars = _page_chars(page, c)
            before = page.render(scale=0.25).to_pil()
            hidden = _hide_text(page, c)
            lines = _lines_from_chars(chars, hidden)
            if hidden and lines:
                after = page.render(scale=0.25).to_pil()
                if before.tobytes() == after.tobytes():
                    lines = []  # текст не удалось убрать с фона — оставим страницу картинкой
        img = page.render(scale=scale, may_draw_forms=True).to_pil().convert("RGB")
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=90)
        buf.seek(0)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(buf, Emu(int(ox)), Emu(int(oy)), Emu(int(w * EMU * kk)), Emu(int(h * EMU * kk)))

        for line in lines:
            real = [ch for ch in line if not ch.get("gen")]
            left = min(ch["l"] for ch in real)
            right = max(ch["r"] for ch in real)
            top = max(ch["t"] for ch in real)
            bottom = min(ch["b"] for ch in real)
            size = max(ch["size"] for ch in real)
            x = ox + left * EMU * kk
            y = oy + (h - top) * EMU * kk
            bw = (right - left + size * 0.6) * EMU * kk * 1.06
            bh = (top - bottom) * EMU * kk
            tb = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(bw)), Emu(int(max(bh, 1))))
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = False
            tf.auto_size = MSO_AUTO_SIZE.NONE
            p = tf.paragraphs[0]
            p.line_spacing = 1.0
            for (font, sz, col, bold, italic), text in _runs(line):
                r = p.add_run()
                r.text = text
                fam, style = _font_family(font)
                sl = style.lower()
                r.font.name = fam
                r.font.size = Pt(max(1, sz * kk * EMU / 12700))
                r.font.bold = bold or any(x in sl for x in ("bold", "black", "heavy", "semi"))
                r.font.italic = italic or "italic" in sl or "oblique" in sl
                r.font.color.rgb = RGBColor(*[max(0, min(255, int(v))) for v in col])
        if lines:
            stats["text_slides"] += 1
        else:
            stats["image_slides"] += 1
        page.close()
        if progress:
            progress(i + 1, n)
    pdf.close()
    prs.save(str(dst))
    return stats


# ---------- Приложение «Конвертер» для macOS (запуск без Терминала) ----------
APP_NAME = "Конвертер"
BUNDLE_VERSION = "2"
BUNDLE = Path.home() / "Applications" / f"{APP_NAME}.app"
DESKTOP = Path.home() / "Desktop"

RUN_SCRIPT = """#!/bin/zsh
# Если конвертер уже запущен — просто открываем страницу
if curl -fs --max-time 1 http://127.0.0.1:8765/status >/dev/null 2>&1; then
  open "http://127.0.0.1:8765/"; exit 0
fi
cd "$HOME/webp-tool" || exit 1
exec .venv/bin/python webp_app.py --app >> "$HOME/webp-tool/log.txt" 2>&1
"""

INFO_PLIST = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE plist PUBLIC "-//Apple//DTD PLIST 1.0//EN" "http://www.apple.com/DTDs/PropertyList-1.0.dtd">
<plist version="1.0"><dict>
  <key>CFBundleName</key><string>{name}</string>
  <key>CFBundleDisplayName</key><string>{name}</string>
  <key>CFBundleIdentifier</key><string>com.mkkatrin.converter</string>
  <key>CFBundleExecutable</key><string>run</string>
  <key>CFBundleIconFile</key><string>icon</string>
  <key>CFBundlePackageType</key><string>APPL</string>
  <key>CFBundleShortVersionString</key><string>{ver}</string>
  <key>CFBundleVersion</key><string>{bver}</string>
  <key>LSUIElement</key><true/>
</dict></plist>
"""


def make_icon(path: Path):
    """Рисует иконку приложения и сохраняет в .icns."""
    from PIL import Image, ImageDraw
    S = 1024
    img = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    grad = Image.new("RGBA", (S, S))
    gd = ImageDraw.Draw(grad)
    for y in range(S):  # градиент сверху вниз: голубой → синий
        t = y / S
        gd.line([(0, y), (S, y)], fill=(int(64 + (10 - 64) * t), int(170 + (90 - 170) * t), 255, 255))
    mask = Image.new("L", (S, S), 0)
    ImageDraw.Draw(mask).rounded_rectangle([100, 100, S - 100, S - 100], radius=185, fill=255)
    img.paste(grad, (0, 0), mask)
    d = ImageDraw.Draw(img)
    # «фотография»: рамка, солнце и горы
    d.rounded_rectangle([250, 290, 774, 734], radius=60, outline=(255, 255, 255, 255), width=44)
    d.ellipse([350, 370, 450, 470], fill=(255, 255, 255, 255))
    d.polygon([(300, 690), (470, 500), (580, 610), (640, 550), (730, 690)], fill=(255, 255, 255, 255))
    img.save(path, sizes=[(16, 16), (32, 32), (64, 64), (128, 128), (256, 256), (512, 512), (1024, 1024)])


def install_bundle(force=False):
    """Создаёт ~/Applications/Конвертер.app и значок на Рабочем столе.
    Возвращает True, если приложение было создано или обновлено сейчас."""
    if sys.platform != "darwin":
        return False
    plist = BUNDLE / "Contents" / "Info.plist"
    if not force and plist.exists() and f"<string>{BUNDLE_VERSION}</string>" in plist.read_text("utf-8"):
        return False
    macos = BUNDLE / "Contents" / "MacOS"
    res = BUNDLE / "Contents" / "Resources"
    macos.mkdir(parents=True, exist_ok=True)
    res.mkdir(parents=True, exist_ok=True)
    run = macos / "run"
    run.write_text(RUN_SCRIPT, "utf-8")
    run.chmod(0o755)
    plist.write_text(INFO_PLIST.format(name=APP_NAME, ver=VERSION, bver=BUNDLE_VERSION), "utf-8")
    try:
        make_icon(res / "icon.icns")
    except Exception as e:
        print("Иконка не создана:", e)
    lsreg = ("/System/Library/Frameworks/CoreServices.framework/Frameworks/"
             "LaunchServices.framework/Support/lsregister")
    for cmd in (["xattr", "-cr", str(BUNDLE)], ["touch", str(BUNDLE)], [lsreg, "-f", str(BUNDLE)]):
        try:
            subprocess.run(cmd, capture_output=True)
        except OSError:
            pass
    # Значок на Рабочем столе вместо старого WebP.command
    try:
        old = DESKTOP / "WebP.command"
        if old.is_file() and "webp-tool" in old.read_text("utf-8", "ignore"):
            old.unlink()
        link = DESKTOP / APP_NAME
        if not link.exists() and not link.is_symlink():
            link.symlink_to(BUNDLE)
    except Exception as e:
        print("Значок на Рабочем столе не создан:", e)
    return True


# ---------- Автовыключение, когда вкладка закрыта ----------
LAST_SEEN = [0.0]
MIGRATED = [False]
IDLE_LIMIT = 180  # секунд без связи со страницей


def watchdog(srv):
    import time
    while True:
        time.sleep(10)
        if LAST_SEEN[0] and time.time() - LAST_SEEN[0] > IDLE_LIMIT:
            srv.shutdown()
            return


def vtuple(v):
    try:
        return tuple(int(x) for x in str(v).split("."))
    except ValueError:
        return (0,)


def fetch(name, timeout=6):
    """Скачать файл из репозитория (через curl — он есть на любом Mac)."""
    r = subprocess.run(["curl", "-fsSL", "--max-time", str(timeout), REPO_RAW + name],
                       capture_output=True)
    if r.returncode != 0:
        raise RuntimeError("нет связи с сервером обновлений")
    return r.stdout


def check_update():
    try:
        info = json.loads(fetch("version.json", 5).decode("utf-8"))
    except Exception:
        return {"available": False, "current": VERSION}
    latest = str(info.get("version", "0"))
    return {"available": vtuple(latest) > vtuple(VERSION), "current": VERSION,
            "latest": latest, "notes": info.get("notes", [])}


def apply_update():
    new = fetch("webp_app.py", 30)
    tmp = APP_FILE.with_suffix(".new")
    tmp.write_bytes(new)
    try:
        compile(new, str(tmp), "exec")  # проверка, что файл целый
    except SyntaxError:
        tmp.unlink()
        raise RuntimeError("скачанный файл повреждён, попробуйте позже")
    if b"VERSION" not in new:
        tmp.unlink()
        raise RuntimeError("неожиданный файл обновления")
    backup = APP_FILE.with_suffix(".py.bak")
    shutil.copy2(APP_FILE, backup)
    os.replace(tmp, APP_FILE)


def restart(port):
    shutil.rmtree(WORK_DIR, ignore_errors=True)
    env = dict(os.environ, WEBP_NO_BROWSER="1", WEBP_PORT=str(port))
    os.execve(sys.executable, [sys.executable, str(APP_FILE)], env)


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _send(self, code, body, ctype):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _json(self, obj, code=200):
        self._send(code, json.dumps(obj, ensure_ascii=False).encode(), "application/json; charset=utf-8")

    def do_GET(self):
        LAST_SEEN[0] = time.time()
        u = urlparse(self.path)
        if u.path == "/ping":
            self._json({"ok": True})
        elif u.path == "/":
            self._send(200, PAGE.encode(), "text/html; charset=utf-8")
        elif u.path == "/status":
            self._json({"ffmpeg": bool(find_ffmpeg()), "out": str(OUT_DIR), "version": VERSION,
                        "migrated": MIGRATED[0], "app": sys.platform == "darwin",
                        "pdf": PDF_STATE["status"]})
        elif u.path.startswith("/pdf/preview/"):
            job = JOBS.get(u.path.rsplit("/", 1)[-1])
            if job and Path(job["res"]).is_file():
                self._send(200, Path(job["res"]).read_bytes(), "application/pdf")
            else:
                self._send(404, b"not found", "text/plain")
        elif u.path == "/update/check":
            self._json(check_update())
        elif u.path.startswith("/out/"):
            f = OUT_DIR / Path(unquote(u.path[5:])).name
            if f.is_file():
                ctype = {".pdf": "application/pdf",
                         ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                         }.get(f.suffix.lower(), "image/webp")
                self._send(200, f.read_bytes(), ctype)
            else:
                self._send(404, b"not found", "text/plain")
        else:
            self._send(404, b"not found", "text/plain")

    def do_POST(self):
        LAST_SEEN[0] = time.time()
        u = urlparse(self.path)
        q = parse_qs(u.query)

        def g(k, d=""):
            return q.get(k, [d])[0]

        length = int(self.headers.get("Content-Length", 0) or 0)

        if u.path == "/update/apply":
            try:
                apply_update()
            except Exception as e:
                return self._json({"ok": False, "error": str(e)})
            self._json({"ok": True})
            port = self.server.server_address[1]
            threading.Timer(0.5, restart, args=(port,)).start()
            return

        if u.path == "/pdf/compress":
            name = Path(g("name", "file.pdf")).name
            try:
                target_mb = max(0.5, float(g("target", "20") or 20))
            except ValueError:
                target_mb = 20.0
            src = WORK_DIR / (os.urandom(6).hex() + ".src.pdf")
            with open(src, "wb") as f:
                left = length
                while left > 0:
                    chunk = self.rfile.read(min(1 << 20, left))
                    if not chunk:
                        break
                    f.write(chunk)
                    left -= len(chunk)
            try:
                wait_pdf_lib()
                return self._json(pdf_compress_job(src, name, target_mb))
            except Exception as e:
                return self._json({"ok": False, "error": str(e)})

        if u.path == "/pdf/pptx":
            name = Path(g("name", "file.pdf")).name
            mode = "image" if g("mode") == "image" else "text"
            src = WORK_DIR / (os.urandom(6).hex() + ".src.pdf")
            with open(src, "wb") as f:
                left = length
                while left > 0:
                    chunk = self.rfile.read(min(1 << 20, left))
                    if not chunk:
                        break
                    f.write(chunk)
                    left -= len(chunk)
            try:
                wait_pdf_lib()
                OUT_DIR.mkdir(parents=True, exist_ok=True)
                dst = unique(OUT_DIR / (Path(name).stem + ".pptx"))
                try:
                    st = pdf_to_pptx(src, dst, mode)
                except Exception as e:
                    if dst.exists():
                        dst.unlink()
                    msg = str(e)
                    if "password" in msg.lower():
                        msg = "PDF защищён паролем — снимите пароль и попробуйте снова"
                    raise RuntimeError(msg)
                return self._json(dict(ok=True, kind="pptx", name=dst.name, url="/out/" + quote(dst.name),
                                       **{"in": length, "out": dst.stat().st_size, "mode": mode}, **st))
            except Exception as e:
                return self._json({"ok": False, "error": str(e)})
            finally:
                src.unlink(missing_ok=True)

        if u.path in ("/pdf/keep", "/pdf/split"):
            jid = g("id")
            if jid not in JOBS:
                return self._json({"ok": False, "error": "задача не найдена — загрузите PDF ещё раз"})
            try:
                if u.path == "/pdf/keep":
                    r = pdf_keep(jid)
                else:
                    r = pdf_split(jid, int(g("parts", "2") or 2))
                return self._json(dict(ok=True, **r))
            except Exception as e:
                return self._json({"ok": False, "error": str(e)})

        if u.path == "/quit":
            self._json({"ok": True})
            threading.Thread(target=self.server.shutdown, daemon=True).start()
            return

        if u.path == "/openfile":
            f = OUT_DIR / Path(g("name")).name
            if f.is_file():
                opener = "open" if sys.platform == "darwin" else ("explorer" if os.name == "nt" else "xdg-open")
                subprocess.Popen([opener, str(f)])
                return self._json({"ok": True})
            return self._json({"ok": False, "error": "файл не найден"})

        if u.path == "/open":
            OUT_DIR.mkdir(parents=True, exist_ok=True)
            opener = "open" if sys.platform == "darwin" else ("explorer" if os.name == "nt" else "xdg-open")
            subprocess.Popen([opener, str(OUT_DIR)])
            return self._json({"ok": True})

        if u.path != "/convert":
            return self._send(404, b"not found", "text/plain")

        name = Path(g("name", "file")).name
        ext = Path(name).suffix.lower()
        if ext not in IMAGE_EXT | VIDEO_EXT:
            self.rfile.read(length)
            return self._json({"ok": False, "error": "Этот формат не поддерживается"})

        quality = max(0, min(100, int(g("q", "80") or 80)))
        lossless = g("lossless") == "1"
        width = int(g("w", "0") or 0) or None
        fps = float(g("fps", "15") or 15)

        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / ("src" + ext)
            with open(src, "wb") as f:
                left = length
                while left > 0:
                    chunk = self.rfile.read(min(1 << 20, left))
                    if not chunk:
                        break
                    f.write(chunk)
                    left -= len(chunk)
            OUT_DIR.mkdir(parents=True, exist_ok=True)
            dst = unique(OUT_DIR / (Path(name).stem + ".webp"))
            try:
                if ext in VIDEO_EXT:
                    convert_video(src, dst, quality, lossless, width, fps, ext == ".gif")
                else:
                    convert_image(src, dst, quality, lossless, width)
            except Exception as e:
                if dst.exists():
                    dst.unlink()
                return self._json({"ok": False, "error": str(e)})

        self._json({"ok": True, "name": dst.name, "in": length,
                    "out": dst.stat().st_size, "url": "/out/" + quote(dst.name)})


PAGE = r"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Конвертер</title>
<style>
:root{
  --bg:#f2f2f4;--panel:#ffffff;--sunk:#f6f6f8;--line:#e2e2e6;--text:#18181b;--muted:#6b6b73;--faint:#9a9aa2;
  --img:#2f6fed;--img-soft:#e8f0fe;--vid:#7c4dff;--vid-soft:#f0ebff;--pdf:#e5484d;--pdf-soft:#fdecec;
  --ok:#16a34a;--warn:#c2410c;--err:#dc2626;--btn:#18181b;--btn-text:#fff;
  --r-lg:16px;--r-md:10px;--r-sm:7px;
}
@media (prefers-color-scheme:dark){:root{
  --bg:#141416;--panel:#1d1d20;--sunk:#252528;--line:#2f2f33;--text:#f2f2f3;--muted:#a1a1a8;--faint:#6f6f76;
  --img:#5b8ff9;--img-soft:#1c2740;--vid:#a07cff;--vid-soft:#2a2240;--pdf:#ff6b6f;--pdf-soft:#3a1f22;
  --ok:#4ade80;--warn:#fb923c;--err:#f87171;--btn:#f2f2f3;--btn-text:#141416;
}}
*{box-sizing:border-box}
html,body{margin:0}
body{background:var(--bg);color:var(--text);font:15px/1.45 -apple-system,BlinkMacSystemFont,"SF Pro Text","Segoe UI",sans-serif;-webkit-font-smoothing:antialiased}
button,input{font:inherit;color:inherit}
:focus-visible{outline:2px solid var(--img);outline-offset:2px}
.wrap{max-width:820px;margin:0 auto;padding:28px 16px 140px}

/* шапка */
.top{display:flex;align-items:center;justify-content:space-between;gap:12px;margin-bottom:20px}
.brand{display:flex;align-items:center;gap:10px}
.logo{width:30px;height:30px;border-radius:8px;background:linear-gradient(#4aa8ff,#1463ff);display:grid;place-items:center}
.logo svg{width:18px;height:18px}
h1{font-size:19px;font-weight:650;letter-spacing:-.01em;margin:0}
.link{background:none;border:0;color:var(--muted);cursor:pointer;padding:6px 8px;border-radius:var(--r-sm);font-size:14px}
.link:hover{color:var(--text);background:var(--sunk)}

/* плашки сверху */
.notice{display:none;margin:0 0 16px;padding:14px 16px;border-radius:var(--r-md);background:var(--panel);border:1px solid var(--line)}
.notice h3{margin:0 0 4px;font-size:15px}
.notice p,.notice ul{margin:0 0 10px;color:var(--muted);font-size:14px}
.notice ul{padding-left:18px}
.notice .row{display:flex;gap:8px;align-items:center;flex-wrap:wrap}
.notice.update{border-color:var(--img)}

/* область загрузки */
.drop{position:relative;border:1.5px dashed var(--line);border-radius:var(--r-lg);background:var(--panel);cursor:pointer;transition:border-color .15s,background .15s;text-align:center}
.drop.empty{padding:88px 24px}
.drop.slim{padding:14px 18px;display:flex;align-items:center;justify-content:center;gap:10px;color:var(--muted)}
.drop.over{border-color:var(--img);background:var(--img-soft)}
.drop .big{font-size:22px;font-weight:650;letter-spacing:-.015em;margin:14px 0 6px}
.drop .small{color:var(--muted);max-width:420px;margin:0 auto}
.kinds{display:flex;justify-content:center;gap:8px;margin-top:22px;flex-wrap:wrap}
.chip{display:inline-flex;align-items:center;gap:6px;font-size:13px;padding:5px 10px;border-radius:99px;background:var(--sunk);color:var(--muted)}
.dot{width:8px;height:8px;border-radius:50%}
.dot.img{background:var(--img)}.dot.vid{background:var(--vid)}.dot.pdf{background:var(--pdf)}
.drop .arrow{width:46px;height:46px;color:var(--faint)}
.drop.slim .arrow{width:18px;height:18px}

/* группы */
.group{margin-top:18px;background:var(--panel);border:1px solid var(--line);border-radius:var(--r-lg);overflow:hidden}
.ghead{display:flex;align-items:center;gap:10px;padding:14px 18px 0}
.ghead h2{margin:0;font-size:16px;font-weight:650}
.ghead .count{color:var(--muted);font-size:14px}
.ghead .bar{width:4px;height:18px;border-radius:2px}
.group.img .bar{background:var(--img)}.group.vid .bar{background:var(--vid)}.group.pdf .bar{background:var(--pdf)}
.settings{display:flex;flex-wrap:wrap;gap:14px 26px;padding:14px 18px 16px;border-bottom:1px solid var(--line)}
.setting label{display:block;font-size:12.5px;color:var(--muted);margin-bottom:6px}
.seg{display:inline-flex;background:var(--sunk);border-radius:var(--r-sm);padding:2px;gap:2px;flex-wrap:wrap}
.seg button{border:0;background:none;padding:5px 11px;border-radius:5px;cursor:pointer;font-size:13.5px;color:var(--muted);white-space:nowrap}
.seg button:hover{color:var(--text)}
.seg button.on{background:var(--panel);color:var(--text);box-shadow:0 0 0 1px var(--line)}
.group.img .seg button.on{color:var(--img)}.group.vid .seg button.on{color:var(--vid)}.group.pdf .seg button.on{color:var(--pdf)}
.custom{width:68px;padding:5px 8px;border:1px solid var(--line);border-radius:var(--r-sm);background:var(--sunk);font-size:13.5px;margin-left:6px}
.hint{flex-basis:100%;font-size:13px;color:var(--muted);margin:-4px 0 0}
.hint.warn{color:var(--warn)}
.hint:empty{display:none}

/* файлы */
.files{list-style:none;margin:0;padding:6px 0}
.file{display:grid;grid-template-columns:44px minmax(0,1fr) auto;gap:4px 14px;align-items:center;padding:8px 18px}
.file + .file{border-top:1px solid var(--line)}
.thumb{width:44px;height:44px;border-radius:8px;background:var(--sunk);object-fit:cover;display:grid;place-items:center;font-size:11px;font-weight:700;overflow:hidden}
.group.pdf .thumb{background:var(--pdf-soft);color:var(--pdf)}
.group.vid .thumb{background:var(--vid-soft);color:var(--vid)}
.fname{font-weight:550;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.fmeta{font-size:13px;color:var(--muted);display:flex;gap:12px;flex-wrap:wrap}
.fmeta .good{color:var(--ok);font-weight:600}
.fmeta .bad{color:var(--err)}
.fmeta .warn{color:var(--warn)}
.side{display:flex;align-items:center;gap:6px}
.x{border:0;background:none;color:var(--faint);cursor:pointer;width:28px;height:28px;border-radius:6px;font-size:18px;line-height:1}
.x:hover{background:var(--sunk);color:var(--text)}
.open{color:var(--img);text-decoration:none;font-weight:600;font-size:14px;padding:5px 8px;border-radius:6px}
.open:hover{background:var(--img-soft)}
.spin{width:18px;height:18px;border:2px solid var(--line);border-top-color:var(--muted);border-radius:50%;animation:s .8s linear infinite}
@keyframes s{to{transform:rotate(360deg)}}
.done-ic{width:20px;height:20px;color:var(--ok)}

/* выбор для PDF */
.choice{grid-column:2 / -1;margin:6px 0 4px;padding:12px 14px;border-radius:var(--r-md);background:var(--pdf-soft);font-size:14px}
.choice p{margin:0 0 10px}
.choice .row{display:flex;gap:8px;flex-wrap:wrap;align-items:center}
.choice small{display:block;color:var(--muted);margin-top:8px;font-size:13px}
.parts{grid-column:2 / -1;display:flex;flex-direction:column;gap:2px;margin:4px 0}
.part{display:flex;justify-content:space-between;align-items:center;gap:10px;font-size:13.5px}
.part span{color:var(--muted)}

/* кнопки */
.btn{border:0;border-radius:var(--r-sm);padding:9px 16px;font-weight:600;cursor:pointer;background:var(--btn);color:var(--btn-text);text-decoration:none;display:inline-block;font-size:14px}
.btn.ghost{background:var(--panel);color:var(--text);box-shadow:inset 0 0 0 1px var(--line)}
.btn:disabled{opacity:.45;cursor:default}

/* нижняя панель */
.dock{position:fixed;left:0;right:0;bottom:0;padding:12px 16px calc(12px + env(safe-area-inset-bottom));background:color-mix(in srgb,var(--bg) 82%,transparent);backdrop-filter:blur(14px);-webkit-backdrop-filter:blur(14px);border-top:1px solid var(--line);display:none}
.dock .in{max-width:820px;margin:0 auto;display:flex;align-items:center;justify-content:space-between;gap:12px}
.dock .sum{color:var(--muted);font-size:14px}
.dock .sum b{color:var(--text);font-weight:600}
.dock .acts{display:flex;gap:8px}
.dock .btn.go{padding:11px 22px;font-size:15px}

.warnline{display:none;margin-top:12px;padding:10px 14px;border-radius:var(--r-md);background:var(--pdf-soft);color:var(--err);font-size:14px}
.unsup{margin-top:12px;font-size:13.5px;color:var(--muted)}
.foot{margin-top:28px;display:flex;justify-content:center;gap:4px;color:var(--faint);font-size:12.5px}
.bye{display:none;text-align:center;padding:120px 16px;color:var(--muted)}
.bye h2{color:var(--text);margin:0 0 6px;font-size:20px}
@media (max-width:560px){.drop.empty{padding:56px 18px}.settings{gap:12px}.file{padding:8px 14px}}
@media (prefers-reduced-motion:reduce){.spin{animation-duration:2s}}
</style></head><body>
<div class="wrap" id="app">
  <div class="top">
    <div class="brand">
      <div class="logo"><svg viewBox="0 0 24 24" fill="none" stroke="#fff" stroke-width="2.2" stroke-linejoin="round"><rect x="3.5" y="5" width="17" height="14" rx="2.5"/><circle cx="8.5" cy="9.5" r="1.6" fill="#fff" stroke="none"/><path d="M5 17l4.5-4.5 3 3 2.5-2.5L19 17"/></svg></div>
      <h1>Конвертер</h1>
    </div>
    <button class="link" id="open">Открыть папку с результатами</button>
  </div>

  <div class="notice" id="note"><h3>Теперь без Терминала</h3><p>Конвертер запускается значком «Конвертер» на Рабочем столе или в Launchpad. Окно Терминала можно закрыть.</p></div>

  <div class="notice update" id="upd">
    <h3 id="updTitle">Доступно обновление</h3>
    <ul id="updNotes"></ul>
    <div class="row"><button class="btn" id="updGo">Обновить</button><button class="btn ghost" id="updLater">Позже</button><span id="updMsg" style="color:var(--muted);font-size:14px"></span></div>
  </div>

  <div class="drop empty" id="drop" tabindex="0" role="button" aria-label="Выбрать файлы">
    <svg class="arrow" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.6" stroke-linecap="round" stroke-linejoin="round"><path d="M12 15V4M7.5 8.5 12 4l4.5 4.5"/><path d="M4 15v3a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2v-3"/></svg>
    <div class="big" id="dropBig">Перетащите файлы сюда</div>
    <div class="small" id="dropSmall">или нажмите, чтобы выбрать. Конвертер сам поймёт, что это, и предложит подходящие настройки.</div>
    <div class="kinds" id="kinds">
      <span class="chip"><span class="dot img"></span>Картинки в WebP</span>
      <span class="chip"><span class="dot vid"></span>Видео и GIF в WebP</span>
      <span class="chip"><span class="dot pdf"></span>Сжатие PDF</span>
    </div>
    <input type="file" id="pick" multiple accept="image/*,video/*,.heic,.heif,.pdf,application/pdf" hidden>
  </div>
  <div class="warnline" id="warn">Видео и GIF не сконвертируются: не установлен ffmpeg. Картинки и PDF работают.</div>
  <div class="unsup" id="unsup"></div>

  <div id="groups"></div>

  <div class="foot"><span id="ver"></span><button class="link" id="quit" style="font-size:12.5px;padding:0 6px">Закрыть конвертер</button></div>
</div>

<div class="bye" id="bye"><h2>Конвертер закрыт</h2>Чтобы снова открыть его, запустите значок «Конвертер».</div>

<div class="dock" id="dock"><div class="in">
  <div class="sum" id="sum"></div>
  <div class="acts"><button class="btn ghost" id="reset">Очистить</button><button class="btn go" id="go">Конвертировать</button></div>
</div></div>

<script>
const $=id=>document.getElementById(id);
const IMG=/\.(jpe?g|png|bmp|tiff?|heic|heif|avif|ico|webp)$/i, VID=/\.(mp4|mov|avi|mkv|webm|m4v|wmv|flv|mpe?g|3gp|gif)$/i, PDF=/\.pdf$/i;
const GROUPS={
  img:{title:'Картинки',settings:{
    q:{label:'Качество',opts:[['lossless','Без потерь'],['90','Высокое'],['80','Баланс'],['65','Компактно']],val:'80'},
    w:{label:'Ширина',opts:[['0','Как есть'],['2560','2560'],['1920','1920'],['1280','1280'],['800','800']],val:'0'}}},
  vid:{title:'Видео и GIF',settings:{
    q:{label:'Качество',opts:[['85','Высокое'],['70','Баланс'],['55','Компактно']],val:'70'},
    w:{label:'Ширина',opts:[['0','Как есть'],['1280','1280'],['720','720'],['480','480']],val:'720'},
    fps:{label:'Плавность',opts:[['10','10 к/с'],['15','15 к/с'],['24','24 к/с'],['30','30 к/с']],val:'15'}}},
  pdf:{title:'PDF',settings:{
    a:{label:'Что сделать',opts:[['compress','Сжать'],['pptx','Превратить в PowerPoint']],val:'compress'},
    t:{label:'Уложиться в',opts:[['5','5 МБ'],['10','10 МБ'],['20','20 МБ'],['50','50 МБ']],val:'20',custom:true,when:['a','compress']},
    m:{label:'Слайды',opts:[['text','С редактируемым текстом'],['image','Точная копия']],val:'text',when:['a','pptx']}}}
};
let items=[],seq=0,busy=false,LATEST='',FFMPEG=true;

function size(n){if(n<1024)return n+' Б';if(n<1048576)return Math.round(n/1024)+' КБ';return (n/1048576).toFixed(1).replace('.',',')+' МБ'}
function esc(s){return String(s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]))}
function plural(n,a,b,c){const m=n%10,h=n%100;return m==1&&h!=11?a:(m>=2&&m<=4&&(h<10||h>=20)?b:c)}
function kindOf(f){if(PDF.test(f.name)||f.type==='application/pdf')return'pdf';if(VID.test(f.name)||f.type.startsWith('video/'))return'vid';if(IMG.test(f.name)||f.type.startsWith('image/'))return'img';return null}
function dur(s){s=Math.round(s);return s<60?s+' с':Math.floor(s/60)+' мин '+(s%60)+' с'}

/* ---------- добавление файлов ---------- */
function add(list){const bad=[];
  for(const f of list){const k=kindOf(f);if(!k){bad.push(f.name);continue}
    items.push({id:++seq,f,k,st:'new'})}
  $('unsup').textContent=bad.length?`Не поддерживается: ${bad.join(', ')}`:'';
  render()}

/* ---------- отрисовка ---------- */
function render(){
  const has=items.length>0;
  const d=$('drop');d.className='drop '+(has?'slim':'empty');
  $('dropBig').style.display=$('dropSmall').style.display=$('kinds').style.display=has?'none':'';
  d.querySelector('.slimtxt')?.remove();
  if(has){const s=document.createElement('span');s.className='slimtxt';s.textContent='Добавить ещё файлы';d.appendChild(s)}
  const box=$('groups');
  for(const k of ['img','vid','pdf']){
    const its=items.filter(i=>i.k===k);let g=$('g-'+k);
    if(!its.length){g?.remove();continue}
    if(!g){g=document.createElement('section');g.id='g-'+k;g.className='group '+k;g.innerHTML=groupHtml(k);
      const next=['img','vid','pdf'].slice(['img','vid','pdf'].indexOf(k)+1).map(x=>$('g-'+x)).find(Boolean);
      box.insertBefore(g,next||null);bindGroup(g,k)}
    g.querySelector('.count').textContent=`${its.length} ${plural(its.length,'файл','файла','файлов')}`;
    const ul=g.querySelector('.files');
    for(const it of its)if(!it.el){it.el=fileEl(it);ul.appendChild(it.el)}
    updateHints(k)}
  updateDock()}

function groupHtml(k){const G=GROUPS[k];
  const sets=Object.entries(G.settings).map(([key,s])=>`<div class="setting" data-key="${key}"${s.when?` data-when="${s.when[0]}=${s.when[1]}"`:''}><label>${s.label}</label><div class="seg" data-k="${key}">${
    s.opts.map(([v,t])=>`<button type="button" data-v="${v}" class="${v===s.val?'on':''}">${t}</button>`).join('')}</div>${
    s.custom?`<input class="custom" type="number" min="1" step="1" placeholder="свой" aria-label="Свой размер в МБ">`:''}</div>`).join('');
  return `<div class="ghead"><span class="bar"></span><h2>${G.title}</h2><span class="count"></span></div>
  <div class="settings">${sets}<p class="hint"></p></div><ul class="files"></ul>`}

function bindGroup(g,k){
  g.querySelectorAll('.seg').forEach(seg=>seg.addEventListener('click',e=>{const b=e.target.closest('button');if(!b)return;
    seg.querySelectorAll('button').forEach(x=>x.classList.toggle('on',x===b));
    GROUPS[k].settings[seg.dataset.k].val=b.dataset.v;const c=g.querySelector('.custom');if(c&&seg.dataset.k==='t')c.value='';applyWhen(g,k);updateHints(k)}));
  applyWhen(g,k);
  const c=g.querySelector('.custom');
  if(c)c.addEventListener('input',()=>{if(c.value){g.querySelectorAll('.seg[data-k=t] button').forEach(x=>x.classList.remove('on'));GROUPS.pdf.settings.t.val=c.value}updateHints(k)})}

function applyWhen(g,k){g.querySelectorAll('[data-when]').forEach(el=>{const [key,v]=el.dataset.when.split('=');el.style.display=GROUPS[k].settings[key].val===v?'':'none'})}
function fileEl(it){const li=document.createElement('li');li.className='file';
  let th;
  if(it.k==='img'&&!/\.(heic|heif|tiff?)$/i.test(it.f.name)){th=`<img class="thumb" src="${URL.createObjectURL(it.f)}" alt="">`}
  else if(it.k==='vid'&&!/\.gif$/i.test(it.f.name)){th=`<video class="thumb" muted preload="metadata" src="${URL.createObjectURL(it.f)}#t=0.1"></video>`}
  else if(it.k==='vid'){th=`<img class="thumb" src="${URL.createObjectURL(it.f)}" alt="">`}
  else th=`<div class="thumb">${it.k==='pdf'?'PDF':esc(it.f.name.split('.').pop().toUpperCase().slice(0,4))}</div>`;
  li.innerHTML=`${th}<div style="min-width:0"><div class="fname">${esc(it.f.name)}</div><div class="fmeta"><span>${size(it.f.size)}</span><span class="extra"></span><span class="st"></span></div></div>
  <div class="side"><button class="x" title="Убрать" aria-label="Убрать файл">×</button></div>`;
  li.querySelector('.x').onclick=()=>{if(it.st==='work')return;items=items.filter(x=>x!==it);li.remove();render()};
  const v=li.querySelector('video');
  if(v)v.addEventListener('loadedmetadata',()=>{it.dur=v.duration;li.querySelector('.extra').textContent=isFinite(v.duration)?dur(v.duration):'';updateHints('vid')});
  return li}

function updateHints(k){const g=$('g-'+k);if(!g)return;const h=g.querySelector('.hint');h.className='hint';
  const its=items.filter(i=>i.k===k&&i.st==='new');
  if(k==='img'){const s=GROUPS.img.settings.q.val;h.textContent=s==='lossless'?'Без потерь подходит для логотипов, иконок и скриншотов. Для фото выберите «Высокое» или «Баланс».':'';}
  if(k==='vid'){const long=its.filter(i=>i.dur>15);
    if(!FFMPEG){h.className='hint warn';h.textContent='Для видео нужен ffmpeg, а он не установлен.'}
    else if(long.length){h.className='hint warn';h.textContent=`${long.length>1?'Есть ролики':'Ролик'} длиннее 15 секунд — анимированный WebP получится тяжёлым. Уменьшите ширину или плавность.`}
    else h.textContent='Звук не сохраняется: WebP — это анимация без звука.'}
  if(k==='pdf'){const S=GROUPS.pdf.settings,t=parseFloat(S.t.val)||20,toPptx=S.a.val==='pptx';
    for(const it of items.filter(i=>i.k==='pdf'&&i.st==='new'))it.el.querySelector('.extra').textContent=!toPptx&&it.f.size<=t*1048576?'уже меньше лимита, только оптимизируем':'';
    h.textContent=!toPptx?'Если для этого придётся сильно ухудшить картинки, Конвертер спросит, оставить один файл или разделить на несколько.'
      :(S.m.val==='text'?'Каждая страница станет слайдом: фон — картинкой, текст — обычными текстовыми блоками, которые можно править. Если в PDF текст нарисован кривыми, такие слайды будут картинками.'
      :'Каждая страница станет слайдом-картинкой: выглядит точно как PDF, но текст не редактируется.')}}

function updateDock(){const pend=items.filter(i=>i.st==='new');const dock=$('dock');
  dock.style.display=items.length?'block':'none';
  const tot=pend.reduce((a,i)=>a+i.f.size,0);
  $('sum').innerHTML=pend.length?`<b>${pend.length} ${plural(pend.length,'файл','файла','файлов')}</b> к конвертации, ${size(tot)}`:(busy?'Конвертирую…':(items.some(i=>i.st==='ask')?'<b>Нужно ваше решение</b> по PDF выше':'Готово. Результаты в папке «Загрузки/Конвертер».'));
  const go=$('go');go.disabled=busy||!pend.length;go.textContent=busy?'Конвертирую…':(pend.length?`Конвертировать ${pend.length} ${plural(pend.length,'файл','файла','файлов')}`:'Конвертировать')}

/* ---------- конвертация ---------- */
function setSide(it,html){it.el.querySelector('.side').innerHTML=html}
const CHECK='<svg class="done-ic" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.4" stroke-linecap="round" stroke-linejoin="round"><path d="M5 12.5l4.5 4.5L19 7"/></svg>';
async function go(){if(busy)return;busy=true;
  const queue=items.filter(i=>i.st==='new');
  queue.forEach(i=>{i.st='wait';i.el.querySelector('.st').textContent='в очереди';setSide(i,'')});
  updateDock();
  for(const it of queue){it.st='work';const st=it.el.querySelector('.st');
    st.textContent=it.k==='pdf'?(GROUPS.pdf.settings.a.val==='pptx'?'делаю презентацию':'сжимаю, большие файлы — до пары минут'):'конвертирую';setSide(it,'<div class="spin"></div>');
    try{let r;
      if(it.k==='pdf'&&GROUPS.pdf.settings.a.val==='pptx'){r=await fetch('/pdf/pptx?'+new URLSearchParams({name:it.f.name,mode:GROUPS.pdf.settings.m.val}),{method:'POST',body:it.f})}
      else if(it.k==='pdf'){r=await fetch('/pdf/compress?'+new URLSearchParams({name:it.f.name,target:GROUPS.pdf.settings.t.val}),{method:'POST',body:it.f})}
      else{const S=GROUPS[it.k].settings;const q=S.q.val;
        const p={name:it.f.name,q:q==='lossless'?'90':q,lossless:q==='lossless'?'1':'0',w:S.w.val,fps:it.k==='vid'?S.fps.val:'15'};
        r=await fetch('/convert?'+new URLSearchParams(p),{method:'POST',body:it.f})}
      const j=await r.json();
      if(!j.ok){it.st='err';st.innerHTML=`<span class="bad">${esc(j.error||'ошибка')}</span>`;setSide(it,'')}
      else if(j.kind==='pdf')pdfDone(it,j);
      else if(j.kind==='pptx')pptxDone(it,j);
      else{it.st='done';const pct=Math.round((1-j.out/j.in)*100);
        st.innerHTML=`→ ${size(j.out)} <span class="${pct>=0?'good':'bad'}">${pct>=0?'−'+pct:'+'+(-pct)}%</span>`;
        setSide(it,`<a class="open" href="/out/${encodeURIComponent(j.name)}" target="_blank">Открыть</a>${CHECK}`)}}
    catch(e){it.st='err';st.innerHTML='<span class="bad">нет связи — конвертер закрыт?</span>';setSide(it,'')}
  }
  busy=false;updateDock()}

function pptxDone(it,j){it.st='done';const st=it.el.querySelector('.st');
  let note=`→ PowerPoint, ${j.slides} ${plural(j.slides,'слайд','слайда','слайдов')}, ${size(j.out)}`;
  if(j.mode==='text'){
    if(j.text_slides===j.slides)note+=', весь текст редактируется';
    else if(j.text_slides===0)note+=`. <span class="warn">Текст в этом PDF нарисован кривыми, поэтому слайды вставлены картинками</span>`;
    else note+=`. <span class="warn">Текст редактируется на ${j.text_slides} из ${j.slides}; на остальных он нарисован кривыми — там картинки</span>`}
  st.innerHTML=note;
  setSide(it,`<button class="open" data-n="${esc(j.name)}" style="border:0;background:none;cursor:pointer">Открыть</button>${CHECK}`);
  it.el.querySelector('.side .open').onclick=e=>fetch('/openfile?'+new URLSearchParams({name:e.target.dataset.n}),{method:'POST'})}
function showParts(it,saved){it.el.querySelector('.parts')?.remove();
  const d=document.createElement('div');d.className='parts';
  d.innerHTML=saved.map(x=>`<div class="part"><div>${esc(x.name)} <span>${size(x.size)}${x.pages?', стр. '+x.pages:''}${x.fits===false?', больше лимита':''}</span></div><a class="open" href="/out/${encodeURIComponent(x.name)}" target="_blank">Открыть</a></div>`).join('');
  it.el.appendChild(d)}

function pdfDone(it,j){const st=it.el.querySelector('.st');it.st='done';
  st.innerHTML=`→ ${size(j.out)}, ${j.pages} ${plural(j.pages,'страница','страницы','страниц')}, ${esc(j.level)}`;
  if(!j.choice){setSide(it,`<a class="open" href="/out/${encodeURIComponent(j.saved[0].name)}" target="_blank">Открыть</a>${CHECK}`);return}
  it.st='ask';setSide(it,'');
  const ch=document.createElement('div');ch.className='choice';
  const why=j.fits?`Чтобы уложиться в ${j.target_mb} МБ, картинки внутри пришлось сильно сжать. Текст останется чётким, а фото и мелкие детали могут стать размытыми.`
    :`Даже при максимальном сжатии получается ${size(j.out)}, а это больше ${j.target_mb} МБ.`;
  ch.innerHTML=`<p>${why}</p><div class="row">
    ${j.parts?`<button class="btn" data-a="split">Разделить на ${j.parts} ${plural(j.parts,'файл','файла','файлов')} с хорошим качеством</button>`:''}
    <button class="btn ghost" data-a="keep">Оставить одним файлом, ${size(j.out)}</button>
    <a class="btn ghost" href="/pdf/preview/${j.id}" target="_blank">Посмотреть результат</a></div>
    <small>Один файл удобнее отправлять, если качество картинок не критично. Разделение сохраняет качество.</small>`;
  it.el.appendChild(ch);
  const act=async(url,label)=>{ch.innerHTML=`<div class="row"><div class="spin"></div><span>${label}</span></div>`;
    try{const r=await (await fetch(url,{method:'POST'})).json();ch.remove();
      if(!r.ok){st.innerHTML=`<span class="bad">${esc(r.error)}</span>`;return}
      if(r.saved.length>1){st.textContent=`разделён на ${r.saved.length} ${plural(r.saved.length,'файл','файла','файлов')}`;showParts(it,r.saved);setSide(it,CHECK)}
      else setSide(it,`<a class="open" href="/out/${encodeURIComponent(r.saved[0].name)}" target="_blank">Открыть</a>${CHECK}`);
      it.st='done';updateDock()}catch(e){ch.innerHTML='<span class="bad">Нет связи — конвертер закрыт?</span>'}};
  ch.querySelector('[data-a=keep]').onclick=()=>act(`/pdf/keep?id=${j.id}`,'Сохраняю…');
  const s=ch.querySelector('[data-a=split]');if(s)s.onclick=()=>act(`/pdf/split?id=${j.id}&parts=${j.parts}`,`Делю на ${j.parts} ${plural(j.parts,'часть','части','частей')} и сжимаю…`)}

/* ---------- события ---------- */
const drop=$('drop'),pick=$('pick');
drop.onclick=()=>pick.click();
drop.onkeydown=e=>{if(e.key==='Enter'||e.key===' '){e.preventDefault();pick.click()}};
pick.onchange=()=>{add(pick.files);pick.value=''};
['dragenter','dragover'].forEach(t=>window.addEventListener(t,e=>{e.preventDefault();drop.classList.add('over')}));
['dragleave','drop'].forEach(t=>window.addEventListener(t,e=>{e.preventDefault();if(t==='drop'||!e.relatedTarget)drop.classList.remove('over')}));
window.addEventListener('drop',e=>{if(e.dataTransfer?.files?.length)add(e.dataTransfer.files)});
$('go').onclick=go;
$('reset').onclick=()=>{if(busy)return;items=items.filter(i=>i.st==='ask');$('groups').innerHTML='';items.forEach(i=>i.el=null);
  if(items.length){/* незавершённые решения по PDF оставляем */}render()};
$('open').onclick=()=>fetch('/open',{method:'POST'});
$('quit').onclick=()=>{fetch('/quit',{method:'POST'}).catch(()=>{});$('app').style.display='none';$('dock').style.display='none';$('bye').style.display='block'};
setInterval(()=>fetch('/ping').catch(()=>{}),20000);

fetch('/status').then(r=>r.json()).then(s=>{$('ver').textContent='Версия '+s.version;FFMPEG=s.ffmpeg;
  if(!s.ffmpeg)$('warn').style.display='block';if(s.migrated)$('note').style.display='block'});
fetch('/update/check').then(r=>r.json()).then(u=>{if(!u.available)return;LATEST=u.latest;
  $('updTitle').textContent=`Доступно обновление ${u.latest}`;
  $('updNotes').innerHTML=(u.notes||[]).map(n=>`<li>${esc(n)}</li>`).join('')||'<li>Улучшения и исправления</li>';
  $('upd').style.display='block'}).catch(()=>{});
$('updLater').onclick=()=>$('upd').style.display='none';
$('updGo').onclick=async()=>{const b=$('updGo');b.disabled=true;$('updLater').disabled=true;$('updMsg').textContent='Скачиваю обновление…';
  try{const j=await (await fetch('/update/apply',{method:'POST'})).json();
    if(!j.ok){$('updMsg').textContent='Не получилось: '+j.error;b.disabled=false;$('updLater').disabled=false;return}
    $('updMsg').textContent='Перезапускаю…';
    for(let i=0;i<40;i++){await new Promise(r=>setTimeout(r,500));try{const st=await (await fetch('/status',{cache:'no-store'})).json();if(st.version===LATEST){location.reload();return}}catch(e){}}
    $('updMsg').textContent='Перезапустите Конвертер вручную'}catch(e){$('updMsg').textContent='Нет связи с конвертером'}};
</script></body></html>
"""


def main():
    if "--install" in sys.argv:
        install_bundle(force=True)
        print(f"Приложение «{APP_NAME}» создано:", BUNDLE)
        return
    threading.Thread(target=ensure_pdf_lib, daemon=True).start()
    try:
        MIGRATED[0] = install_bundle()
    except Exception as e:
        print("Не удалось создать приложение:", e)
    srv = None
    pref = int(os.environ.get("WEBP_PORT", "0") or 0)
    ports = ([pref] if pref else []) + [8765, 8766, 8767, 0]
    for port in ports:
        try:
            srv = ThreadingHTTPServer(("127.0.0.1", port), Handler)
            break
        except OSError:
            continue
    url = f"http://127.0.0.1:{srv.server_address[1]}/"
    print(f"{APP_NAME} {VERSION} запущен:", url)
    print("Готовые файлы:", OUT_DIR)
    LAST_SEEN[0] = time.time()
    threading.Thread(target=watchdog, args=(srv,), daemon=True).start()
    if not os.environ.get("WEBP_NO_BROWSER"):
        threading.Timer(0.6, lambda: webbrowser.open(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        shutil.rmtree(WORK_DIR, ignore_errors=True)


if __name__ == "__main__":
    main()
